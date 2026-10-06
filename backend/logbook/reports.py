"""Report builder: filters + optional grouping + chosen columns -> table with totals.

Only whitelisted field, metric and dimension names reach SQL; values are always bound parameters.
"""
import csv
import io

import openpyxl
from openpyxl.styles import Font, PatternFill

from .fields import DIMENSIONS, FIELD_KIND, FIELD_LABEL, FIELD_NAMES, METRICS, ROLES, minutes_to_hours


class ReportError(ValueError):
    pass


def _where(filters):
    where, args = [], []
    f = filters or {}
    if f.get("date_from"):
        where.append("f.date >= ?"); args.append(f["date_from"])
    if f.get("date_to"):
        where.append("f.date <= ?"); args.append(f["date_to"])
    for key, column in (("type_codes", "f.type_code"), ("registrations", "f.registration"),
                        ("categories", "t.category")):
        values = [v for v in f.get(key) or [] if v]
        if values:
            where.append(f"{column} IN ({', '.join('?' * len(values))})"); args += values
    kind = f.get("kind")
    if kind == "aircraft":
        where.append("f.is_sim = 0")
    elif kind == "sim":
        where.append("f.is_sim = 1")
    elif kind == "flight_time":
        where.append("f.flight_time > 0")
    roles = [r for r in f.get("roles") or [] if r in ROLES + ["instructor", "examiner"]]
    if roles:
        where.append("(" + " OR ".join(f"f.{r} > 0" for r in roles) + ")")
    conditions = [c for c in f.get("conditions") or [] if c in ("night", "ifr", "nvg", "xc", "multi_pilot", "ldg_ship")]
    for c in conditions:
        where.append("(f.ifr_actual + f.ifr_sim) > 0" if c == "ifr" else f"f.{c} > 0")
    if f.get("pf_pm") in ("PF", "PM"):
        where.append("f.pf_pm = ?"); args.append(f["pf_pm"])
    if f.get("place"):
        where.append("(f.dep = ? OR f.arr = ?)"); args += [f["place"].upper()] * 2
    if f.get("text"):
        like = f"%{f['text']}%"
        where.append("(f.remarks LIKE ? OR f.tags LIKE ? OR f.name_pic LIKE ? OR f.name_copilot LIKE ? "
                     "OR f.name_instructor LIKE ? OR f.name_student LIKE ?)")
        args += [like] * 6
    return (f"WHERE {' AND '.join(where)}" if where else ""), args


def _format(kind, value):
    return minutes_to_hours(value) if kind == "duration" else (value or 0)


def run(conn, spec):
    """spec: {filters, mode: summary|detail, group_by: [dim], columns: [metric or field]}"""
    mode = spec.get("mode", "summary")
    columns = spec.get("columns") or ["count", "flight_time"]
    where, args = _where(spec.get("filters"))
    base = "FROM flight f LEFT JOIN aircraft_type t ON t.code = f.type_code"

    if mode == "detail":
        bad = [c for c in columns if c not in FIELD_NAMES and c not in METRICS]
        if bad:
            raise ReportError(f"Unknown columns: {bad}")
        cols = [c for c in columns if c in FIELD_NAMES or c in ("ifr", "landings")]
        select = ", ".join("f.ifr_actual + f.ifr_sim" if c == "ifr" else
                           "f.ldg_day + f.ldg_night" if c == "landings" else f"f.{c}" for c in cols)
        rows = conn.execute(f"SELECT {select} {base} {where} ORDER BY f.date, f.id", args).fetchall()
        kinds = [METRICS[c][1] if c in ("ifr", "landings") else FIELD_KIND[c] for c in cols]
        headers = [{"key": c, "label": METRICS[c][0] if c in ("ifr", "landings") else FIELD_LABEL[c], "kind": k}
                   for c, k in zip(cols, kinds)]
        out_rows = [[_format(k, v) if k in ("duration", "count") else v for k, v in zip(kinds, r)] for r in rows]
        totals = [_format(k, sum((r[i] or 0) for r in rows)) if k in ("duration", "count") else None
                  for i, k in enumerate(kinds)]
        if totals:
            totals[0] = f"{len(rows)} entries" if totals[0] is None else totals[0]
        return {"headers": headers, "rows": out_rows, "totals": totals}

    group_by = spec.get("group_by") or []
    bad = [g for g in group_by if g not in DIMENSIONS] + [c for c in columns if c not in METRICS]
    if bad:
        raise ReportError(f"Unknown group or column: {bad}")
    dims = [DIMENSIONS[g][1] for g in group_by]
    aggs = [METRICS[c][2] for c in columns]
    select = ", ".join([f"{d} AS g{i}" for i, d in enumerate(dims)] + aggs)
    group = f"GROUP BY {', '.join(f'g{i}' for i in range(len(dims)))} ORDER BY {', '.join(f'g{i}' for i in range(len(dims)))}" if dims else ""
    rows = conn.execute(f"SELECT {select} {base} {where} {group}", args).fetchall()
    total = conn.execute(f"SELECT {', '.join(aggs)} {base} {where}", args).fetchone()
    kinds = [METRICS[c][1] for c in columns]
    headers = [{"key": g, "label": DIMENSIONS[g][0], "kind": "text"} for g in group_by] + \
              [{"key": c, "label": METRICS[c][0], "kind": k} for c, k in zip(columns, kinds)]
    n = len(dims)
    out_rows = [list(r[:n]) + [_format(k, v) for k, v in zip(kinds, r[n:])] for r in rows]
    totals = (["Total"] + [None] * (n - 1) if n else []) + [_format(k, v) for k, v in zip(kinds, total)]
    return {"headers": headers, "rows": out_rows, "totals": totals}


def to_csv(result):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h["label"] for h in result["headers"]])
    w.writerows(result["rows"])
    w.writerow(["" if v is None else v for v in result["totals"]])
    return buf.getvalue().encode()


def to_xlsx(result, title="Logbook report"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([title])
    ws["A1"].font = Font(name="Arial", bold=True, size=13)
    ws.append([h["label"] for h in result["headers"]])
    for c in ws[2]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="1F3864")
    for r in result["rows"]:
        ws.append(r)
    ws.append(result["totals"])
    for c in ws[ws.max_row]:
        c.font = Font(name="Arial", bold=True)
    for i, h in enumerate(result["headers"]):
        col = ws.cell(2, i + 1).column_letter
        ws.column_dimensions[col].width = 34 if h["key"] == "remarks" else max(10, len(h["label"]) + 3)
        if h["kind"] == "duration":
            for cell in ws[col][2:]:
                cell.number_format = "0.0"
    for row in ws.iter_rows(min_row=3):
        for c in row:
            if not c.font.bold:
                c.font = Font(name="Arial")
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
